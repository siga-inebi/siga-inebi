"""
Notas por curso desde Excel, cursos visibles del docente e informe de calificaciones.

- Cursos: dirección ve todos los del plan; un docente, solo sus asignaciones.
- Plantilla: estudiantes reales con su código, una columna por unidad activa.
- Importación: por código, todo o nada, reglas de captura, consultas constantes.
- Informe: promedios enteros con redondeo mitad hacia arriba, filtros, permiso
  ``grade.report`` y consultas constantes.
"""

import gzip
import io
import json
import zipfile
from datetime import timedelta
from types import SimpleNamespace

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.academics.models import CurriculumPlan, TeachingAssignment
from apps.audit.models import AuditEvent
from apps.enrolments.services import create_enrolment
from apps.evaluation.models import Grade
from apps.evaluation.workbook import HEADER_ROW, METADATA_SHEET
from tests.factories.academic import AcademicCycleFactory, SectionFactory, SubjectFactory
from tests.factories.evaluation import EvaluationUnitFactory, grant_scoped_permission, open_unit
from tests.factories.identity import UserFactory
from tests.factories.people import PersonFactory
from tests.factories.students import StudentFactory

pytestmark = [pytest.mark.api, pytest.mark.django_db]


def _enrol(cycle, section, *, first_name="Ana", last_name="Zeta", code=None):
    extra = {"student_code": code} if code else {}
    student = StudentFactory(
        person=PersonFactory(first_name=first_name, last_name=last_name), **extra
    )
    return create_enrolment(
        student=student, academic_cycle=cycle, grade=section.grade, section=section
    )


def _closed_window_unit(cycle, number):
    """A unit of the past: its capture window ended before today."""
    today = timezone.localdate()
    start = today - timedelta(days=200 - number * 40)
    return EvaluationUnitFactory(
        academic_cycle=cycle,
        number=number,
        name=f"Unidad {number}",
        starts_on=start,
        ends_on=start + timedelta(days=20),
        capture_starts_on=start,
        capture_ends_on=start + timedelta(days=25),
    )


def _world(user, institution, *, students=3):
    """
    A cycle with one section, Matemática in its plan, a past unit (window
    closed) and a current one (window open), the user as the teacher of the
    current assignment holding grade.write scoped to it.
    """
    cycle = AcademicCycleFactory(institution=institution)
    section = SectionFactory(academic_cycle=cycle)
    subject = SubjectFactory(institution=institution, name="Matemática")
    CurriculumPlan.objects.create(academic_cycle=cycle, grade=section.grade, subject=subject)
    past = _closed_window_unit(cycle, 1)
    current = open_unit(cycle, number=2, name="Unidad 2")
    enrolments = [
        _enrol(cycle, section, first_name=f"Est{index}", last_name=f"Ape{index:02d}")
        for index in range(students)
    ]
    TeachingAssignment.objects.create(
        academic_cycle=cycle,
        section=section,
        subject=subject,
        teacher=user.person,
        starts_on=cycle.starts_on,
    )
    grant_scoped_permission(user, "grade_write")
    return SimpleNamespace(
        cycle=cycle,
        section=section,
        subject=subject,
        past=past,
        current=current,
        enrolments=enrolments,
    )


def _url(name, world):
    return reverse(name, kwargs={"cycle_public_id": str(world.cycle.public_id)})


def _course(world):
    return {"section_id": str(world.section.public_id), "subject_id": str(world.subject.public_id)}


def _download(client, world):
    response = client.get(_url("course-grade-workbook", world), _course(world))
    assert response.status_code == 200, response.content[:500]
    return load_workbook(io.BytesIO(response.content))


def _upload(client, world, workbook, *, name="notas.xlsx", course=None):
    buffer = io.BytesIO()
    workbook.save(buffer)
    payload = {
        "file": SimpleUploadedFile(name, buffer.getvalue()),
        **(course or _course(world)),
    }
    return client.post(_url("course-grade-workbook-import", world), payload)


def _row_of(sheet, code):
    for row in range(HEADER_ROW + 1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == code:
            return row
    raise AssertionError(f"code {code} not in template")


class TestGradeCourses:
    def test_teacher_sees_only_their_assigned_courses(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        other_section = SectionFactory(academic_cycle=world.cycle, grade=world.section.grade)
        other_subject = SubjectFactory(institution=institution, name="Idioma")
        CurriculumPlan.objects.create(
            academic_cycle=world.cycle, grade=world.section.grade, subject=other_subject
        )
        TeachingAssignment.objects.create(
            academic_cycle=world.cycle,
            section=other_section,
            subject=other_subject,
            teacher=PersonFactory(),
            starts_on=world.cycle.starts_on,
        )

        response = auth_client.get(_url("grade-courses", world))

        assert response.status_code == 200
        courses = response.json()
        assert [(c["section"]["public_id"], c["subject"]["name"]) for c in courses] == [
            (str(world.section.public_id), "Matemática")
        ]
        assert courses[0]["teacher_name"] == (
            f"{auth_client.user.person.first_name} {auth_client.user.person.last_name}"
        )

    def test_institution_wide_writer_sees_every_planned_course(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        other_subject = SubjectFactory(institution=institution, name="Idioma")
        CurriculumPlan.objects.create(
            academic_cycle=world.cycle, grade=world.section.grade, subject=other_subject
        )
        grant_scoped_permission(auth_client.user, "grade_write", institution=institution)

        courses = auth_client.get(_url("grade-courses", world)).json()

        assert sorted(course["subject"]["name"] for course in courses) == ["Idioma", "Matemática"]
        idioma = next(course for course in courses if course["subject"]["name"] == "Idioma")
        assert idioma["teacher_name"] is None

    def test_account_without_grade_write_gets_an_empty_list(self, client, institution):
        user = UserFactory()
        client.force_login(user)
        cycle = AcademicCycleFactory(institution=institution)

        response = client.get(
            reverse("grade-courses", kwargs={"cycle_public_id": str(cycle.public_id)})
        )

        assert response.status_code == 200
        assert response.json() == []


class TestCourseWorkbookExport:
    def test_template_has_real_students_codes_and_every_active_unit(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        first = world.enrolments[0]
        Grade.objects.create(
            enrolment=first, subject=world.subject, evaluation_unit=world.past, value=77
        )

        workbook = _download(auth_client, world)
        sheet = workbook["Notas"]

        headers = [sheet.cell(row=HEADER_ROW, column=col).value for col in range(1, 5)]
        assert headers == ["Código", "Estudiante", "Unidad 1 (cerrada)", "Unidad 2"]
        row = _row_of(sheet, first.student.student_code)
        assert sheet.cell(row=row, column=2).value == "Est0 Ape00"
        assert sheet.cell(row=row, column=3).value == 77
        assert sheet.cell(row=row, column=4).value is None
        codes = [
            sheet.cell(row=r, column=1).value
            for r in range(HEADER_ROW + 1, HEADER_ROW + 1 + len(world.enrolments))
        ]
        assert sorted(codes) == sorted(e.student.student_code for e in world.enrolments)
        # Código y nombre bloqueados; la unidad abierta editable, la cerrada no.
        assert sheet.protection.sheet is True
        assert sheet.cell(row=row, column=1).protection.locked is True
        assert sheet.cell(row=row, column=3).protection.locked is True
        assert sheet.cell(row=row, column=4).protection.locked is False
        assert workbook[METADATA_SHEET].sheet_state == "veryHidden"

    def test_unit_columns_follow_the_cycle_configuration(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        open_unit(
            world.cycle,
            number=3,
            name="Unidad 3",
            starts_on=world.current.ends_on + timedelta(days=1),
            ends_on=world.current.ends_on + timedelta(days=10),
        )

        sheet = _download(auth_client, world)["Notas"]

        assert sheet.cell(row=HEADER_ROW, column=5).value == "Unidad 3"

    def test_teacher_of_another_course_is_denied(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        foreign = SubjectFactory(institution=institution)

        response = auth_client.get(
            _url("course-grade-workbook", world),
            {"section_id": str(world.section.public_id), "subject_id": str(foreign.public_id)},
        )

        assert response.status_code == 403
        assert AuditEvent.objects.filter(action="evaluation.grade_write_denied").exists()


class TestCourseWorkbookImport:
    def test_imports_by_code_and_ignores_name_differences(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        workbook = _download(auth_client, world)
        sheet = workbook["Notas"]
        sheet.protection.sheet = False
        first, second = world.enrolments[:2]
        row = _row_of(sheet, first.student.student_code)
        sheet.cell(row=row, column=2, value="ÉST0 ápe00 (mal escrito)")
        sheet.cell(row=row, column=4, value=88)
        sheet.cell(row=_row_of(sheet, second.student.student_code), column=4, value="91")

        response = _upload(auth_client, world, workbook)

        assert response.status_code == 200, response.json()
        assert response.json() == {"created": 2, "updated": 0, "unchanged": 0}
        values = dict(
            Grade.objects.filter(evaluation_unit=world.current).values_list("enrolment_id", "value")
        )
        assert values == {first.pk: 88, second.pk: 91}
        assert AuditEvent.objects.filter(action="evaluation.grade_registered").count() == 2
        assert AuditEvent.objects.filter(action="evaluation.grades_workbook_imported").exists()

    def test_correction_is_audited_with_before_and_after(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        first = world.enrolments[0]
        grade = Grade.objects.create(
            enrolment=first, subject=world.subject, evaluation_unit=world.current, value=50
        )
        workbook = _download(auth_client, world)
        sheet = workbook["Notas"]
        sheet.cell(row=_row_of(sheet, first.student.student_code), column=4, value=65)

        response = _upload(auth_client, world, workbook)

        assert response.json() == {"created": 0, "updated": 1, "unchanged": 0}
        grade.refresh_from_db()
        assert grade.value == 65
        event = AuditEvent.objects.get(action="evaluation.grade_updated")
        assert event.context["changes"] == {"value": {"before": 50, "after": 65}}
        assert event.context["source"] == "workbook"

    def test_reimporting_the_same_file_changes_nothing(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        first = world.enrolments[0]
        Grade.objects.create(
            enrolment=first, subject=world.subject, evaluation_unit=world.past, value=70
        )
        Grade.objects.create(
            enrolment=first, subject=world.subject, evaluation_unit=world.current, value=80
        )

        response = _upload(auth_client, world, _download(auth_client, world))

        # La columna de la unidad cerrada trae la nota registrada: no es un error.
        assert response.status_code == 200, response.json()
        assert response.json() == {"created": 0, "updated": 0, "unchanged": 2}

    def test_empty_cell_never_erases_a_registered_grade(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        first = world.enrolments[0]
        Grade.objects.create(
            enrolment=first, subject=world.subject, evaluation_unit=world.current, value=80
        )
        workbook = _download(auth_client, world)
        sheet = workbook["Notas"]
        sheet.cell(row=_row_of(sheet, first.student.student_code), column=4, value=None)

        response = _upload(auth_client, world, workbook)

        assert response.status_code == 200
        assert Grade.objects.get(enrolment=first, evaluation_unit=world.current).value == 80

    def test_any_invalid_cell_rejects_the_whole_file(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        workbook = _download(auth_client, world)
        sheet = workbook["Notas"]
        sheet.protection.sheet = False
        good, bad = world.enrolments[:2]
        sheet.cell(row=_row_of(sheet, good.student.student_code), column=4, value=90)
        bad_row = _row_of(sheet, bad.student.student_code)
        sheet.cell(row=bad_row, column=4, value=120)
        unknown_row = sheet.max_row + 1
        sheet.cell(row=unknown_row, column=1, value="NO-EXISTE")
        sheet.cell(row=unknown_row, column=4, value=70)

        response = _upload(auth_client, world, workbook)

        assert response.status_code == 400
        errors = response.json()["errors"]
        assert {"row": bad_row, "student_code": bad.student.student_code} == {
            k: errors[0][k] for k in ("row", "student_code")
        }
        assert errors[0]["unit"] == "Unidad 2"
        assert "entre 0 y 100" in errors[0]["message"]
        assert errors[1]["row"] == unknown_row
        assert "NO-EXISTE" in errors[1]["message"]
        assert not Grade.objects.exists()

    def test_change_in_a_closed_window_unit_is_rejected(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        workbook = _download(auth_client, world)
        sheet = workbook["Notas"]
        sheet.protection.sheet = False
        first = world.enrolments[0]
        row = _row_of(sheet, first.student.student_code)
        sheet.cell(row=row, column=3, value=60)

        response = _upload(auth_client, world, workbook)

        assert response.status_code == 400
        error = response.json()["errors"][0]
        assert error["row"] == row
        assert error["unit"] == "Unidad 1"
        assert "no admite captura" in error["message"]
        assert not Grade.objects.exists()

    def test_duplicated_code_is_rejected(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        workbook = _download(auth_client, world)
        sheet = workbook["Notas"]
        sheet.protection.sheet = False
        code = world.enrolments[0].student.student_code
        extra = sheet.max_row + 1
        sheet.cell(row=extra, column=1, value=code)

        response = _upload(auth_client, world, workbook)

        assert response.status_code == 400
        assert "más de una vez" in response.json()["errors"][0]["message"]

    def test_template_of_another_course_is_rejected(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        workbook = _download(auth_client, world)
        other_subject = SubjectFactory(institution=institution, name="Idioma")
        CurriculumPlan.objects.create(
            academic_cycle=world.cycle, grade=world.section.grade, subject=other_subject
        )
        TeachingAssignment.objects.create(
            academic_cycle=world.cycle,
            section=world.section,
            subject=other_subject,
            teacher=auth_client.user.person,
            starts_on=world.cycle.starts_on,
        )

        response = _upload(
            auth_client,
            world,
            workbook,
            course={
                "section_id": str(world.section.public_id),
                "subject_id": str(other_subject.public_id),
            },
        )

        assert response.status_code == 400
        assert "otro curso" in json.dumps(response.json(), ensure_ascii=False)

    def test_files_that_are_not_the_template_are_rejected(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        url = _url("course-grade-workbook-import", world)

        not_xlsx = auth_client.post(
            url, {"file": SimpleUploadedFile("notas.csv", b"a,b"), **_course(world)}
        )
        broken = auth_client.post(
            url, {"file": SimpleUploadedFile("notas.xlsx", b"no es zip"), **_course(world)}
        )
        foreign = io.BytesIO()
        from openpyxl import Workbook

        Workbook().save(foreign)
        plain = auth_client.post(
            url,
            {"file": SimpleUploadedFile("notas.xlsx", foreign.getvalue()), **_course(world)},
        )

        assert not_xlsx.status_code == 400
        assert broken.status_code == 400
        assert plain.status_code == 400
        assert "plantilla de notas de SIGA" in json.dumps(plain.json(), ensure_ascii=False)

    def test_zip_bomb_is_rejected_before_parsing(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/huge.xml", b"0" * (30 * 1024 * 1024))

        response = auth_client.post(
            _url("course-grade-workbook-import", world),
            {"file": SimpleUploadedFile("notas.xlsx", bomb.getvalue()), **_course(world)},
        )

        assert response.status_code == 400
        assert "demasiado grande" in json.dumps(response.json(), ensure_ascii=False)

    def test_import_cost_does_not_grow_with_rows(self, auth_client, institution):
        world = _world(auth_client.user, institution, students=3)

        def fill_and_upload(value):
            workbook = _download(auth_client, world)
            sheet = workbook["Notas"]
            for enrolment in world.enrolments:
                sheet.cell(
                    row=_row_of(sheet, enrolment.student.student_code), column=4
                ).value = value
            with CaptureQueriesContext(connection) as queries:
                response = _upload(auth_client, world, workbook)
            assert response.status_code == 200, response.json()
            return len(queries)

        few = fill_and_upload(70)
        for index in range(30):
            world.enrolments.append(_enrol(world.cycle, world.section, first_name=f"Extra{index}"))
        Grade.objects.all().delete()
        many = fill_and_upload(70)

        assert many == few


class TestGradeReport:
    def _report_world(self, user, institution):
        world = _world(user, institution, students=2)
        idioma = SubjectFactory(institution=institution, name="Idioma")
        CurriculumPlan.objects.create(
            academic_cycle=world.cycle, grade=world.section.grade, subject=idioma
        )
        world.idioma = idioma
        grant_scoped_permission(user, "grade_report", institution=institution)
        return world

    def test_averages_round_half_up_and_skip_ungraded_units(self, auth_client, institution):
        world = self._report_world(auth_client.user, institution)
        ana = world.enrolments[0]
        for subject, unit, value in (
            (world.subject, world.past, 60),
            (world.subject, world.current, 61),
            (world.idioma, world.past, 90),
        ):
            Grade.objects.create(enrolment=ana, subject=subject, evaluation_unit=unit, value=value)

        response = auth_client.get(
            _url("grade-report", world), {"student_id": str(ana.student.public_id)}
        )

        assert response.status_code == 200, response.content[:300]
        data = response.json()
        assert [unit["name"] for unit in data["units"]] == ["Unidad 1", "Unidad 2"]
        assert data["institution"]["name"] == institution.name
        (student,) = data["students"]
        assert student["student_code"] == ana.student.student_code
        subjects = {subject["name"]: subject for subject in student["subjects"]}
        # Idioma (alfabético) antes que Matemática; 60.5 -> 61, nunca 60.
        assert [subject["name"] for subject in student["subjects"]] == ["Idioma", "Matemática"]
        assert subjects["Matemática"] == {"name": "Matemática", "grades": [60, 61], "average": 61}
        assert subjects["Idioma"] == {"name": "Idioma", "grades": [90, None], "average": 90}
        assert student["unit_averages"] == [75, 61]
        # (60 + 61 + 90) / 3 = 70.33 -> 70
        assert student["general_average"] == 70

    def test_student_without_grades_has_null_averages(self, auth_client, institution):
        world = self._report_world(auth_client.user, institution)

        data = auth_client.get(_url("grade-report", world)).json()

        assert len(data["students"]) == 2
        student = data["students"][0]
        assert student["general_average"] is None
        assert student["unit_averages"] == [None, None]
        assert all(subject["average"] is None for subject in student["subjects"])

    def test_filters_by_grade_and_section(self, auth_client, institution):
        world = self._report_world(auth_client.user, institution)
        sibling = SectionFactory(academic_cycle=world.cycle, grade=world.section.grade)
        _enrol(world.cycle, sibling, first_name="Otra")
        url = _url("grade-report", world)
        grade_id = str(world.section.grade.public_id)

        by_grade = auth_client.get(url, {"grade_id": grade_id}).json()
        by_section = auth_client.get(
            url, {"grade_id": grade_id, "section_id": str(sibling.public_id)}
        ).json()
        section_without_grade = auth_client.get(url, {"section_id": str(sibling.public_id)})

        assert len(by_grade["students"]) == 3
        assert [s["student_name"] for s in by_section["students"]] == ["Otra Zeta"]
        assert section_without_grade.status_code == 400

    def test_requires_grade_report_and_records_the_read(self, auth_client, institution):
        world = _world(auth_client.user, institution)
        url = _url("grade-report", world)

        denied = auth_client.get(url)
        grant_scoped_permission(auth_client.user, "grade_report", institution=institution)
        allowed = auth_client.get(url)

        assert denied.status_code == 403
        assert AuditEvent.objects.filter(action="evaluation.grade_report_denied").exists()
        assert allowed.status_code == 200
        event = AuditEvent.objects.get(action="evaluation.grade_report_generated")
        assert event.context["student_count"] == 3

    def test_teacher_grade_write_does_not_allow_printing(self, auth_client, institution):
        world = _world(auth_client.user, institution)

        assert auth_client.get(_url("grade-report", world)).status_code == 403

    def test_response_is_gzip_compressed(self, auth_client, institution):
        world = self._report_world(auth_client.user, institution)

        response = auth_client.get(_url("grade-report", world), HTTP_ACCEPT_ENCODING="gzip")

        assert response.status_code == 200
        assert response["Content-Encoding"] == "gzip"
        assert json.loads(gzip.decompress(response.content))["students"]

    def test_report_cost_does_not_grow_with_students(self, auth_client, institution):
        world = self._report_world(auth_client.user, institution)
        url = _url("grade-report", world)
        with CaptureQueriesContext(connection) as few:
            assert auth_client.get(url).status_code == 200

        for index in range(25):
            enrolment = _enrol(world.cycle, world.section, first_name=f"Más{index}")
            Grade.objects.create(
                enrolment=enrolment, subject=world.subject, evaluation_unit=world.past, value=80
            )
        with CaptureQueriesContext(connection) as many:
            response = auth_client.get(url)

        assert len(response.json()["students"]) == 27
        assert len(many) == len(few)
