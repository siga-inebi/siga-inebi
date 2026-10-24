"""
Plantilla Excel de notas por curso: transporte, sin reglas de negocio.

Este módulo solo traduce entre un libro ``.xlsx`` y estructuras de Python. Qué
estudiantes van, qué unidades admiten captura y qué significa cada celda al
importar lo decide ``apps.evaluation.services`` (ADR-0008).

La plantilla lleva una hoja muy oculta (``_siga``) con los identificadores del
ciclo, la sección, el curso y cada columna de unidad. Así una importación sabe
exactamente a qué curso y a qué unidad pertenece cada celda aunque alguien
renombre una columna, y rechaza el archivo de otro curso antes de leer notas.
"""

import io
import zipfile
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.common.exceptions import DomainError

TEMPLATE_KIND = "siga-course-grades"
TEMPLATE_VERSION = "1"
GRADES_SHEET = "Notas"
METADATA_SHEET = "_siga"

#: Fila de encabezados y primeras columnas. Fijas en la versión 1 de la
#: plantilla y repetidas en los metadatos, que es lo que lee la importación.
HEADER_ROW = 6
CODE_COLUMN = 1
NAME_COLUMN = 2
FIRST_UNIT_COLUMN = 3

#: Topes defensivos de un archivo subido. Una sección real son ~40 filas y unos
#: pocos KB; nada legítimo se acerca a estos números.
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 200
MAX_DATA_ROWS = 1000

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_BRAND = "1F4E79"
_HEADER_FILL = PatternFill("solid", fgColor=_BRAND)
_LOCKED_FILL = PatternFill("solid", fgColor="E7E6E6")
_ZEBRA_FILL = PatternFill("solid", fgColor="F3F7FB")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build_course_workbook(data):
    """
    Bytes of the template for one course.

    ``data`` keys: ``institution_name``, ``cycle_name``, ``section_label``,
    ``subject_name``, ``teacher_name`` (or None), ``generated_at`` (datetime),
    ``cycle_id``, ``section_id``, ``subject_id`` (str UUIDs), ``units``
    (``[{"id", "label", "editable", "lock_reason"}]``) and ``rows``
    (``[{"student_code", "student_name", "values": [int|None per unit]}]``).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = GRADES_SHEET
    units = data["units"]
    last_column = FIRST_UNIT_COLUMN + len(units) - 1 if units else NAME_COLUMN

    _write_heading(sheet, data, last_column)
    _write_headers(sheet, units)
    _write_rows(sheet, units, data["rows"])
    _write_validation(sheet, units, len(data["rows"]))

    sheet.column_dimensions[get_column_letter(CODE_COLUMN)].width = 16
    sheet.column_dimensions[get_column_letter(NAME_COLUMN)].width = 38
    for offset in range(len(units)):
        sheet.column_dimensions[get_column_letter(FIRST_UNIT_COLUMN + offset)].width = 16
    sheet.freeze_panes = sheet.cell(row=HEADER_ROW + 1, column=FIRST_UNIT_COLUMN)
    # Sin contraseña: evita editar por accidente códigos, nombres y columnas
    # cerradas, no es un control de seguridad (el backend vuelve a validar todo).
    sheet.protection.sheet = True
    sheet.protection.formatColumns = False
    sheet.protection.selectLockedCells = False
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1

    _write_metadata(workbook, data)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_heading(sheet, data, last_column):
    span = max(last_column, NAME_COLUMN)
    lines = [
        (data["institution_name"], Font(bold=True, size=14, color=_BRAND)),
        (
            f"Notas de {data['subject_name']} · {data['section_label']} · {data['cycle_name']}",
            Font(bold=True, size=12),
        ),
        (
            f"Docente: {data['teacher_name'] or 'sin asignación vigente'} · Generado el "
            f"{data['generated_at']:%d/%m/%Y %H:%M}",
            Font(size=10, color="595959"),
        ),
        (
            "Escriba notas enteras de 0 a 100 en las columnas de cada unidad. No cambie los "
            "códigos: con ellos se identifica a cada estudiante. Una celda vacía no modifica "
            "la nota registrada. Las columnas grises no admiten captura hoy.",
            Font(size=10, italic=True, color="595959"),
        ),
    ]
    for index, (text, font) in enumerate(lines, start=1):
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=span)
        cell = sheet.cell(row=index, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[4].height = 42


def _write_headers(sheet, units):
    headers = ["Código", "Estudiante"] + [unit["label"] for unit in units]
    for column, text in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    for offset, unit in enumerate(units):
        if not unit["editable"]:
            header = sheet.cell(row=HEADER_ROW, column=FIRST_UNIT_COLUMN + offset)
            header.value = f"{unit['label']} (cerrada)"
    sheet.row_dimensions[HEADER_ROW].height = 32


def _write_rows(sheet, units, rows):
    for index, row in enumerate(rows):
        excel_row = HEADER_ROW + 1 + index
        zebra = _ZEBRA_FILL if index % 2 else None
        code = sheet.cell(row=excel_row, column=CODE_COLUMN, value=str(row["student_code"]))
        code.number_format = "@"
        name = sheet.cell(row=excel_row, column=NAME_COLUMN, value=row["student_name"])
        for cell in (code, name):
            cell.border = _BORDER
            cell.protection = Protection(locked=True)
            if zebra:
                cell.fill = zebra
        for offset, unit in enumerate(units):
            cell = sheet.cell(
                row=excel_row, column=FIRST_UNIT_COLUMN + offset, value=row["values"][offset]
            )
            cell.alignment = Alignment(horizontal="center")
            cell.border = _BORDER
            cell.number_format = "0"
            if unit["editable"]:
                cell.protection = Protection(locked=False)
                if zebra:
                    cell.fill = zebra
            else:
                cell.protection = Protection(locked=True)
                cell.fill = _LOCKED_FILL


def _write_validation(sheet, units, row_count):
    if not row_count:
        return
    first_row = HEADER_ROW + 1
    last_row = HEADER_ROW + row_count
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1="0",
        formula2="100",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Nota fuera de escala",
        error="Escriba un número entero de 0 a 100.",
    )
    for offset, unit in enumerate(units):
        if unit["editable"]:
            letter = get_column_letter(FIRST_UNIT_COLUMN + offset)
            validation.add(f"{letter}{first_row}:{letter}{last_row}")
    if validation.sqref:
        sheet.add_data_validation(validation)


def _write_metadata(workbook, data):
    meta = workbook.create_sheet(METADATA_SHEET)
    meta.sheet_state = "veryHidden"
    values = [
        ("kind", TEMPLATE_KIND),
        ("version", TEMPLATE_VERSION),
        ("cycle_id", data["cycle_id"]),
        ("section_id", data["section_id"]),
        ("subject_id", data["subject_id"]),
        ("header_row", HEADER_ROW),
        ("first_unit_column", FIRST_UNIT_COLUMN),
    ]
    for index, (key, value) in enumerate(values, start=1):
        meta.cell(row=index, column=1, value=key)
        meta.cell(row=index, column=2, value=value)
    meta.cell(row=len(values) + 1, column=1, value="unit_ids")
    for offset, unit in enumerate(data["units"]):
        meta.cell(row=len(values) + 1, column=2 + offset, value=unit["id"])


def read_course_workbook(uploaded):
    """
    Parse an uploaded template into ``{"metadata": {...}, "unit_ids": [...],
    "rows": [{"row": int, "student_code": str, "student_name": str,
    "cells": [raw per unit column]}]}``.

    Raw cell values are returned untouched (int, float, str or None): turning
    them into grades is a domain rule. Raises ``DomainError`` when the file is
    not a SIGA course template or exceeds the defensive limits.
    """
    content = _read_bounded(uploaded)
    _check_archive(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except Exception as exc:  # openpyxl raises many types for a broken file
        raise DomainError(
            "No se pudo leer el archivo. Suba la plantilla .xlsx descargada desde SIGA."
        ) from exc
    try:
        metadata, unit_ids = _read_metadata(workbook)
        if GRADES_SHEET not in workbook.sheetnames:
            raise DomainError("La plantilla no tiene la hoja «Notas».")
        rows = _read_rows(workbook[GRADES_SHEET], metadata, len(unit_ids))
    finally:
        workbook.close()
    return {"metadata": metadata, "unit_ids": unit_ids, "rows": rows}


def _read_bounded(uploaded):
    if uploaded is None:
        raise DomainError("Adjunte la plantilla de notas en formato .xlsx.")
    name = (getattr(uploaded, "name", "") or "").lower()
    if not name.endswith(".xlsx"):
        raise DomainError("El archivo debe ser la plantilla .xlsx descargada desde SIGA.")
    size = getattr(uploaded, "size", None)
    if size is not None and size > MAX_UPLOAD_BYTES:
        raise DomainError("El archivo supera 2 MB; no parece una plantilla de notas.")
    content = uploaded.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise DomainError("El archivo supera 2 MB; no parece una plantilla de notas.")
    return content


def _check_archive(content):
    """Reject zip bombs before openpyxl inflates anything."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise DomainError(
            "El archivo no es un .xlsx válido. Suba la plantilla descargada desde SIGA."
        ) from exc
    if len(members) > MAX_ARCHIVE_MEMBERS or (
        sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES
    ):
        raise DomainError("El archivo es demasiado grande para ser una plantilla de notas.")


def _read_metadata(workbook):
    if METADATA_SHEET not in workbook.sheetnames:
        raise DomainError(
            "El archivo no es una plantilla de notas de SIGA. Descargue la plantilla del curso "
            "y escriba las notas en ella."
        )
    metadata = {}
    unit_ids = []
    for row in workbook[METADATA_SHEET].iter_rows(max_row=20, values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0])
        if key == "unit_ids":
            unit_ids = [str(value) for value in row[1:] if value not in (None, "")]
        else:
            metadata[key] = row[1] if len(row) > 1 else None
    if metadata.get("kind") != TEMPLATE_KIND:
        raise DomainError("El archivo no es una plantilla de notas de SIGA.")
    if str(metadata.get("version")) != TEMPLATE_VERSION:
        raise DomainError(
            "La plantilla es de una versión anterior. Descargue una plantilla nueva del curso."
        )
    for key in ("cycle_id", "section_id", "subject_id"):
        metadata[key] = str(metadata.get(key) or "")
    return metadata, unit_ids


def _read_rows(sheet, metadata, unit_count):
    header_row = int(metadata.get("header_row") or HEADER_ROW)
    first_unit = int(metadata.get("first_unit_column") or FIRST_UNIT_COLUMN)
    last_column = max(first_unit + unit_count - 1, NAME_COLUMN)
    rows = []
    for number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, max_col=last_column, values_only=True),
        start=header_row + 1,
    ):
        values = list(values) + [None] * (last_column - len(values))
        code = _as_code(values[CODE_COLUMN - 1])
        name = values[NAME_COLUMN - 1]
        cells = values[first_unit - 1 : first_unit - 1 + unit_count]
        if not code and name in (None, "") and all(cell in (None, "") for cell in cells):
            continue
        rows.append(
            {
                "row": number,
                "student_code": code,
                "student_name": "" if name is None else str(name).strip(),
                "cells": cells,
            }
        )
        if len(rows) > MAX_DATA_ROWS:
            raise DomainError(
                f"La plantilla tiene más de {MAX_DATA_ROWS} filas; no corresponde a un curso."
            )
    return rows


def _as_code(value):
    """Student code as text; Excel may have turned a numeric code into a float."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()
